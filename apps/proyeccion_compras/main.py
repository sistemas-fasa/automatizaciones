#!/usr/bin/env python3
"""
Proyección de Compras Automatizada
-----------------------------------
Toma todos los artículos activos con stock, identifica su proveedor principal,
calcula proyección basada en 12 meses de histórico y envía reporte por email.

Uso:
    python main.py                    # corrida completa
    python main.py --dry-run          # solo log, sin email
    python main.py --no-email         # genera Excel pero no envía email
"""

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

from config import cfg
from calculator import calcular_proyecciones, resumen_a_dataframe
from mailer import send_email

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
    stream=sys.stdout,
)
logger = logging.getLogger("main")


def exportar_excel(resumenes, output_dir: str) -> Path:
    """Genera Excel con una hoja por proveedor + resumen general."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath = out_dir / f"proyeccion_compras_{timestamp}.xlsx"

    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    header_align = Alignment(horizontal="center", vertical="center")
    num_align = Alignment(horizontal="right")

    headers = [
        "Proveedor",
        "Articulo",
        "Detalle",
        "Unidad",
        "Stock Actual",
        "Prom. Mensual",
        "Proyeccion",
        "Saldo Pedidos",
        "Pedido Sugerido",
    ]

    # --- Hoja Resumen ---
    ws_sum = wb.active
    ws_sum.title = "Resumen"
    sum_headers = [
        "Proveedor",
        "Nombre",
        "Artículos",
        "Total Stock",
        "Total Proyección",
        "Total Pedido Sugerido",
    ]
    for col, h in enumerate(sum_headers, 1):
        cell = ws_sum.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for i, prov in enumerate(resumenes, 2):
        ws_sum.cell(row=i, column=1, value=prov.codigo)
        ws_sum.cell(row=i, column=2, value=prov.nombre)
        ws_sum.cell(row=i, column=3, value=prov.cantidad_articulos)
        ws_sum.cell(row=i, column=4, value=prov.total_stock).alignment = num_align
        ws_sum.cell(row=i, column=5, value=prov.total_proyeccion).alignment = num_align
        ws_sum.cell(
            row=i, column=6, value=prov.total_pedido_sugerido
        ).alignment = num_align

    # Anchos columna resumen
    widths = [10, 35, 10, 14, 16, 18]
    for col, w in enumerate(widths, 1):
        ws_sum.column_dimensions[get_column_letter(col)].width = w

    # --- Hoja por proveedor ---
    for prov in resumenes:
        safe_name = prov.codigo
        ws = wb.create_sheet(safe_name)

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for i, art in enumerate(prov.articulos, 2):
            ws.cell(row=i, column=1, value=f"{prov.codigo} - {prov.nombre}")
            ws.cell(row=i, column=2, value=art.articulo)
            ws.cell(row=i, column=3, value=art.detalle)
            ws.cell(row=i, column=4, value=art.unidad)
            ws.cell(row=i, column=5, value=art.stock_actual).alignment = num_align
            ws.cell(row=i, column=6, value=art.promedio_mensual).alignment = num_align
            ws.cell(row=i, column=7, value=art.proyeccion_bruta).alignment = num_align
            ws.cell(row=i, column=8, value=art.saldo_pedidos).alignment = num_align
            ws.cell(row=i, column=9, value=art.pedido_sugerido).alignment = num_align

        col_widths = [35, 10, 40, 8, 14, 14, 14, 14, 16]
        for col, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w

    wb.save(filepath)
    logger.info("Excel generado: %s", filepath)
    return filepath


def main():
    parser = argparse.ArgumentParser(description="Proyección de Compras Automatizada")
    parser.add_argument(
        "--dry-run", action="store_true", help="Solo log, sin email ni Excel"
    )
    parser.add_argument(
        "--no-email", action="store_true", help="Genera Excel pero no envía email"
    )
    args = parser.parse_args()

    logger.info("=== INICIO Proyección de Compras ===")
    logger.info(
        "Meses a proyectar: %d | Empresa ID: %d", cfg.proy_meses, cfg.empresa_id
    )

    if args.dry_run:
        logger.info("Modo DRY RUN: calculando sin generar archivos ni email")
        resumenes = calcular_proyecciones()
        logger.info("Proveedores con proyección: %d", len(resumenes))
        for prov in resumenes:
            logger.info(
                "  Prov %s - %s: %d artículos, pedido sugerido total: %.2f",
                prov.codigo,
                prov.nombre,
                prov.cantidad_articulos,
                prov.total_pedido_sugerido,
            )
        logger.info("=== FIN DRY RUN ===")
        return

    resumenes = calcular_proyecciones()

    if not resumenes:
        logger.warning("No se generaron proyecciones. Saliendo.")
        return

    excel_path = exportar_excel(resumenes, cfg.excel_output_dir)

    if not args.no_email:
        send_email(resumenes, excel_path)
    else:
        logger.info("Modo --no-email: email omitido")

    logger.info("=== FIN Proyección de Compras ===")


if __name__ == "__main__":
    main()
